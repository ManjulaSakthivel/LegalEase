from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

from utils.text_utils import sanitize_text


BASE_DIR = Path(__file__).resolve().parents[1]
LOGO_PATH = BASE_DIR / "assets" / "logo.png"


def _ensure_logo() -> Path | None:
    # Logo PNG is optional. SVG is included for source/design purposes.
    return LOGO_PATH if LOGO_PATH.exists() else None


def format_docx(text: str, doc_type: str) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    logo = _ensure_logo()
    if logo:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(logo), width=Inches(1.0))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for block in re.split(r"\n\s*\n", text.strip()):
        block = block.strip()
        if not block:
            continue

        first_line = block.splitlines()[0].strip()
        is_heading = bool(
            re.match(r"^\d+[\.\)]\s+", first_line)
            or first_line.isupper() and len(first_line) < 100
        )

        if is_heading:
            p = doc.add_paragraph()
            r = p.add_run(block)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
        else:
            p = doc.add_paragraph(block)
            p.paragraph_format.space_after = Pt(6)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase - AI-assisted draft. Review before legal use.")

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


class LegalEasePDF(FPDF):
    def __init__(self, doc_type: str) -> None:
        super().__init__()
        self.doc_type = doc_type
        self.set_auto_page_break(auto=True, margin=18)

    def header(self) -> None:
        self.set_font("Helvetica", "B", 10)
        self.cell(0, 8, self.doc_type.upper(), align="C")
        self.ln(6)

    def footer(self) -> None:
        self.set_y(-14)
        self.set_font("Helvetica", "I", 8)
        self.cell(
            0,
            8,
            f"LegalEase - AI-assisted draft | Page {self.page_no()}",
            align="C",
        )


def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = LegalEasePDF(doc_type)
    pdf.set_title(doc_type)
    pdf.add_page()
    pdf.set_font("Times", size=11)

    safe_text = sanitize_text(text)
    for block in re.split(r"\n\s*\n", safe_text):
        block = block.strip()
        if not block:
            continue

        first_line = block.splitlines()[0]
        heading = bool(
            re.match(r"^\d+[\.\)]\s+", first_line)
            or first_line.isupper() and len(first_line) < 100
        )

        if heading:
            pdf.set_font("Times", "B", 12)
            pdf.multi_cell(0, 7, block)
            pdf.ln(2)
            pdf.set_font("Times", size=11)
        else:
            pdf.multi_cell(0, 6, block)
            pdf.ln(2)

    return bytes(pdf.output())
